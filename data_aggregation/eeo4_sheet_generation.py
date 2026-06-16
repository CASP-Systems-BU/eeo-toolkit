"""
This script post-processes differentially private EEO-4 contingency tables.
It expects the outputs of eeo4_dp.py (WFRG_all_hires.csv, WFRG_new_hires.csv,
JRG.csv, TFG.csv, TFR.csv) and produces:
1. Raw and zeroized Excel workbooks for initial review.
2. Adjusted 3-way side tables with non-negative counts (cvxpy MIP).
3. An adjusted 4-way WFRG table (Work_Salary × Government_Function × Race × Sex).
4. A final adjusted Excel workbook for publication.
"""

import os
import pandas as pd
import numpy as np
import cvxpy as cp
from itertools import combinations
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo

# Directory containing WFRG_all_hires.csv, WFRG_new_hires.csv, JRG.csv, TFG.csv, TFR.csv
input_dir = "/home/eolwd/data/eeo4_csv"

# Directory for Excel workbook output
output_dir = "/home/eolwd/data/eeo4_csv"

# Column renames: DP script output names → publication-format names
RENAME_DICT = {
    'Job Category':        'Job_Category',
    'Work Type':           'Work_Type',
    'Salary Range Groups': 'Salary_Range',
    'Government Function': 'Government_Function',
    'Gov Type':            'Government_Type',
    'Race':                'Race_Ethnicity',
    'Gender':              'Sex',
}


# =============================================================================
# 3-way table adjustment helpers
# =============================================================================

def _effective_delta(orig, corr):
    """Asymmetric cost: reaching 0 from a negative cell is free; excess above 0 is charged."""
    return pd.Series(np.where(
        orig >= 0,
        (corr - orig).abs(),
        corr.clip(lower=0),
    ), index=orig.index)


def _print_3way_report(csv_path, df_orig, df_corr, all_vars, count_col,
                       max_cell_change, max_1way, max_2way, solver_used, obj_value):
    orig_neg  = (df_orig[count_col] < 0).sum()
    corr_neg  = (df_corr[count_col] < 0).sum()
    orig_tot  = df_orig[count_col].sum()
    corr_tot  = df_corr[count_col].sum()
    pct_shift = (corr_tot - orig_tot) / orig_tot * 100 if orig_tot else 0

    print(f"\n{'='*60}")
    print(f"  {csv_path}")
    print(f"  Solver : {solver_used}  |  obj = {obj_value:.1f}")
    print(f"  Neg cells : {orig_neg} → {corr_neg}  |  "
          f"total : {orig_tot:.1f} → {corr_tot} ({pct_shift:+.3f}%)")
    print(f"{'='*60}")

    merged    = df_orig.merge(df_corr, on=all_vars, suffixes=('_orig', '_corr'))
    orig_vals = merged[f'{count_col}_orig']
    corr_vals = merged[f'{count_col}_corr']
    merged['eff_change'] = _effective_delta(orig_vals, corr_vals)
    merged['raw_change'] = corr_vals - orig_vals

    changed = (merged[orig_vals != corr_vals]
               .assign(_abs=lambda d: d['raw_change'].abs())
               .sort_values('_abs', ascending=False)
               .drop(columns='_abs')
               .copy())

    print(f"  Cells changed: {len(changed)}")
    if not changed.empty:
        print(changed[all_vars + [f'{count_col}_orig', f'{count_col}_corr',
                                  'raw_change', 'eff_change']]
              .rename(columns={f'{count_col}_orig': 'original',
                               f'{count_col}_corr': 'corrected'})
              .to_string(index=False))

    print()
    for v in all_vars:
        delta  = (df_corr.groupby(v)[count_col].sum() - df_orig.groupby(v)[count_col].sum()).abs()
        status = "✓" if max_1way is None or delta.max() <= max_1way else f"✗ max exceeded ({delta.max()})"
        print(f"  1-way [{v}]: max |Δ| = {delta.max():.0f}  "
              f"(limit={max_1way if max_1way is not None else '—'})  {status}")

    for v0, v1 in combinations(all_vars, 2):
        delta  = (df_corr.groupby([v0, v1])[count_col].sum() - df_orig.groupby([v0, v1])[count_col].sum()).abs()
        status = "✓" if max_2way is None or delta.max() <= max_2way else f"✗ max exceeded ({delta.max()})"
        print(f"  2-way [{v0} × {v1}]: max |Δ| = {delta.max():.0f}  "
              f"(limit={max_2way if max_2way is not None else '—'})  {status}")

    violations = (merged['eff_change'] > max_cell_change).sum()
    print(f"  Cell change (limit={max_cell_change}): "
          f"{'✓' if not violations else f'✗ {violations} violation(s)'}")


def adjust_3way_table(
    csv_path,
    count_col='Count',
    use_l1=True,
    max_cell_change=10,
    max_1way_marginal_change=None,
    max_2way_marginal_change=None,
    output_csv=None,
):
    """
    Correct a noisy 3-way DP contingency table so every cell is non-negative.

    Objective: minimize L1 (or L2) error on the 3-way cells.
    Constraints:
      * All cells >= 0
      * Per-cell budget (asymmetric for originally-negative cells):
          - positive original: |corrected - original| <= max_cell_change
          - negative original: corrected <= max_cell_change (reaching 0 is free)
      * Each 1-way marginal entry may change by at most max_1way_marginal_change
      * Each 2-way marginal entry may change by at most max_2way_marginal_change
    """
    df_3way  = pd.read_csv(csv_path)
    all_vars = [c for c in df_3way.columns if c != count_col]

    if len(all_vars) != 3:
        raise ValueError(
            f"{csv_path}: expected 3 categorical columns + '{count_col}', "
            f"found {len(all_vars)}: {all_vars}"
        )

    categories = {v: sorted(df_3way[v].unique()) for v in all_vars}
    n_cats     = {v: len(c) for v, c in categories.items()}
    index_map  = {v: {cat: i for i, cat in enumerate(cats)} for v, cats in categories.items()}
    shape      = tuple(n_cats[v] for v in all_vars)

    Y = np.zeros(shape)
    for _, row in df_3way.iterrows():
        idx    = tuple(index_map[v][row[v]] for v in all_vars)
        Y[idx] = row[count_col]

    X = cp.Variable(shape, integer=True)
    objective = cp.Minimize(cp.sum(cp.abs(X - Y)) if use_l1 else cp.sum_squares(X - Y))

    upper_bound = np.maximum(Y, 0) + max_cell_change
    lower_bound = Y - max_cell_change
    constraints = [X >= 0, X >= lower_bound, X <= upper_bound]

    if max_1way_marginal_change is not None:
        for i, v in enumerate(all_vars):
            other_axes = tuple(j for j in range(3) if j != i)
            X_1way = cp.sum(X, axis=other_axes)
            Y_1way = np.sum(Y, axis=other_axes)
            constraints += [
                X_1way - Y_1way <= max_1way_marginal_change,
                Y_1way - X_1way <= max_1way_marginal_change,
            ]

    if max_2way_marginal_change is not None:
        for v_pair in combinations(range(3), 2):
            remaining_axis = next(j for j in range(3) if j not in v_pair)
            X_2way = cp.sum(X, axis=remaining_axis)
            Y_2way = np.sum(Y, axis=remaining_axis)
            constraints += [
                X_2way - Y_2way <= max_2way_marginal_change,
                Y_2way - X_2way <= max_2way_marginal_change,
            ]

    problem     = cp.Problem(objective, constraints)
    solved      = False
    solver_used = None

    for name, solver in [("CBC", cp.CBC), ("SCIP", cp.SCIP), ("GLPK_MI", cp.GLPK_MI)]:
        try:
            problem.solve(solver=solver, verbose=False)
            if problem.status in ("optimal", "optimal_inaccurate"):
                solved      = True
                solver_used = name
                break
        except Exception:
            pass

    if not solved:
        raise RuntimeError(
            f"No MIP solver succeeded for {csv_path} (status: {problem.status})."
        )

    X_int = np.round(X.value).astype(int)

    rows = []
    for idx in np.ndindex(shape):
        row            = {v: list(categories[v])[idx[i]] for i, v in enumerate(all_vars)}
        row[count_col] = X_int[idx]
        rows.append(row)
    df_corrected = pd.DataFrame(rows)

    _print_3way_report(
        csv_path, df_3way, df_corrected, all_vars, count_col,
        max_cell_change, max_1way_marginal_change, max_2way_marginal_change,
        solver_used, problem.value,
    )

    if output_csv:
        df_corrected.to_csv(output_csv, index=False)
        print(f"  → saved to {output_csv}")

    return df_corrected


# =============================================================================
# 4-way table adjustment helpers
# =============================================================================

def _marginal_report(df_orig, df_corr, group_cols, count_col, limit, label):
    """Print a one-line marginal shift summary with asymmetric effective delta."""
    orig_m = (df_orig.groupby(group_cols)[count_col].sum() if group_cols
              else pd.Series([df_orig[count_col].sum()], index=[0]))
    corr_m = (df_corr.groupby(group_cols)[count_col].sum() if group_cols
              else pd.Series([df_corr[count_col].sum()], index=[0]))
    eff = np.where(orig_m >= 0, (corr_m - orig_m).abs(), corr_m.clip(lower=0))
    max_eff = eff.max()
    status = ("✓" if limit is None or max_eff <= limit
              else f"✗ max exceeded ({max_eff:.0f})")
    print(f"  {label}: max |Δ| = {max_eff:.0f}  "
          f"(limit={limit if limit is not None else '—'})  {status}")


def _print_4way_report(csv_path, df_orig, df_corr, all_vars, count_col,
                       max_cell_change, max_1way, max_2way, max_3way, max_0way,
                       solver_used, obj_value):
    orig_neg  = (df_orig[count_col] < 0).sum()
    corr_neg  = (df_corr[count_col] < 0).sum()
    orig_tot  = df_orig[count_col].sum()
    corr_tot  = df_corr[count_col].sum()
    pct_shift = (corr_tot - orig_tot) / orig_tot * 100 if orig_tot else 0

    print(f"\n{'='*60}")
    print(f"  {csv_path}")
    print(f"  Solver : {solver_used}  |  obj = {obj_value:.1f}")
    print(f"  Neg cells : {orig_neg} → {corr_neg}  |  "
          f"total : {orig_tot:.1f} → {corr_tot} ({pct_shift:+.3f}%)")
    print(f"{'='*60}")

    merged    = df_orig.merge(df_corr, on=all_vars, suffixes=('_orig', '_corr'))
    orig_vals = merged[f'{count_col}_orig']
    corr_vals = merged[f'{count_col}_corr']
    merged['eff_change'] = _effective_delta(orig_vals, corr_vals)
    merged['raw_change'] = corr_vals - orig_vals

    changed = (merged[orig_vals != corr_vals]
               .assign(_abs=lambda d: d['raw_change'].abs())
               .sort_values('_abs', ascending=False)
               .drop(columns='_abs')
               .copy())

    print(f"  Cells changed: {len(changed)}")
    if not changed.empty:
        print(changed[all_vars + [f'{count_col}_orig', f'{count_col}_corr',
                                  'raw_change', 'eff_change']]
              .rename(columns={f'{count_col}_orig': 'original',
                               f'{count_col}_corr': 'corrected'})
              .to_string(index=False))

    violations = (merged['eff_change'] > max_cell_change).sum()
    print(f"\n  Cell change (limit={max_cell_change}): "
          f"{'✓' if not violations else f'✗ {violations} violation(s)'}")

    print()
    _marginal_report(df_orig, df_corr, [],           count_col, max_0way, "0-way [grand total]")
    for v in all_vars:
        _marginal_report(df_orig, df_corr, [v],      count_col, max_1way, f"1-way [{v}]")
    for v0, v1 in combinations(all_vars, 2):
        _marginal_report(df_orig, df_corr, [v0, v1], count_col, max_2way, f"2-way [{v0} × {v1}]")
    for trio in combinations(all_vars, 3):
        _marginal_report(df_orig, df_corr, list(trio), count_col, max_3way,
                         f"3-way [{' × '.join(trio)}]")


def adjust_4way_table(
    csv_path,
    count_col='Count',
    use_l1=True,
    max_cell_change=10,
    max_1way_marginal_change=None,
    max_2way_marginal_change=None,
    max_3way_marginal_change=None,
    max_0way_marginal_change=None,
    output_csv=None,
):
    """
    Correct a noisy 4-way DP contingency table so every cell is non-negative.

    Objective: minimize L1 (or L2) error on the 4-way cells.
    Constraints:
      * All cells >= 0
      * Per-cell budget (asymmetric for originally-negative cells):
          - positive original: |corrected - original| <= max_cell_change
          - negative original: corrected <= max_cell_change (reaching 0 is free)
      * Each k-way marginal (k=0,1,2,3) may change by at most max_kway_marginal_change,
        with the same asymmetric treatment: reaching 0 from negative is free.
    """
    df_4way  = pd.read_csv(csv_path)
    all_vars = [c for c in df_4way.columns if c != count_col]

    if len(all_vars) != 4:
        raise ValueError(
            f"{csv_path}: expected 4 categorical columns + '{count_col}', "
            f"found {len(all_vars)}: {all_vars}"
        )

    categories = {v: sorted(df_4way[v].unique()) for v in all_vars}
    n_cats     = {v: len(c) for v, c in categories.items()}
    index_map  = {v: {cat: i for i, cat in enumerate(cats)} for v, cats in categories.items()}
    shape      = tuple(n_cats[v] for v in all_vars)

    Y = np.zeros(shape)
    for _, row in df_4way.iterrows():
        idx    = tuple(index_map[v][row[v]] for v in all_vars)
        Y[idx] = row[count_col]

    X = cp.Variable(shape, integer=True)
    objective = cp.Minimize(cp.sum(cp.abs(X - Y)) if use_l1 else cp.sum_squares(X - Y))

    upper_bound = np.maximum(Y, 0) + max_cell_change
    lower_bound = Y - max_cell_change
    constraints = [X >= 0, X <= upper_bound, X >= lower_bound]

    def _add_marginal_constraints(X_marg, Y_marg, limit):
        constraints.append(X_marg <= np.maximum(Y_marg, 0) + limit)
        constraints.append(X_marg >= Y_marg - limit)

    if max_3way_marginal_change is not None:
        for trio in combinations(range(4), 3):
            remaining = next(j for j in range(4) if j not in trio)
            _add_marginal_constraints(
                cp.sum(X, axis=remaining), np.sum(Y, axis=remaining),
                max_3way_marginal_change,
            )

    if max_2way_marginal_change is not None:
        for pair in combinations(range(4), 2):
            remaining = tuple(j for j in range(4) if j not in pair)
            _add_marginal_constraints(
                cp.sum(X, axis=remaining), np.sum(Y, axis=remaining),
                max_2way_marginal_change,
            )

    if max_1way_marginal_change is not None:
        for i in range(4):
            other = tuple(j for j in range(4) if j != i)
            _add_marginal_constraints(
                cp.sum(X, axis=other), np.sum(Y, axis=other),
                max_1way_marginal_change,
            )

    if max_0way_marginal_change is not None:
        Y_total = float(np.sum(Y))
        constraints.append(cp.sum(X) <= max(Y_total, 0) + max_0way_marginal_change)
        constraints.append(cp.sum(X) >= Y_total - max_0way_marginal_change)

    problem     = cp.Problem(objective, constraints)
    solved      = False
    solver_used = None

    for name, solver in [("CBC", cp.CBC), ("SCIP", cp.SCIP), ("GLPK_MI", cp.GLPK_MI)]:
        try:
            problem.solve(solver=solver, verbose=False)
            if problem.status in ("optimal", "optimal_inaccurate"):
                solved      = True
                solver_used = name
                break
        except Exception:
            pass

    if not solved:
        raise RuntimeError(
            f"No MIP solver succeeded for {csv_path} (status: {problem.status})."
        )

    X_int = np.round(X.value).astype(int)

    rows = []
    for idx in np.ndindex(shape):
        row            = {v: list(categories[v])[idx[i]] for i, v in enumerate(all_vars)}
        row[count_col] = X_int[idx]
        rows.append(row)
    df_corrected = pd.DataFrame(rows)

    _print_4way_report(
        csv_path, df_4way, df_corrected, all_vars, count_col,
        max_cell_change,
        max_1way_marginal_change, max_2way_marginal_change,
        max_3way_marginal_change, max_0way_marginal_change,
        solver_used, problem.value,
    )

    if output_csv:
        df_corrected.to_csv(output_csv, index=False)
        print(f"  → saved to {output_csv}")

    return df_corrected


# =============================================================================
# Excel generation helpers
# =============================================================================

def write_tables_with_titles(filename, dataframes_dict):
    """Write DataFrames to Excel, each on its own sheet with a title row and a formatted table."""
    with pd.ExcelWriter(filename, engine='openpyxl') as writer:
        for sheet_name, df in dataframes_dict.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=1)

    wb = load_workbook(filename)

    for sheet_name, df in dataframes_dict.items():
        ws = wb[sheet_name]

        ws['A1'] = sheet_name
        ws['A1'].font = Font(bold=True, size=12)

        max_row = ws.max_row
        max_col = ws.max_column
        tab_ref = f"A2:{get_column_letter(max_col)}{max_row}"

        tab = Table(displayName=sheet_name.replace('-', '_').replace(' ', '_'), ref=tab_ref)
        tab.tableStyleInfo = TableStyleInfo(
            name=None, showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(tab)

        for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            for cell in row:
                cell.font = Font(
                    name=cell.font.name, size=12,
                    bold=cell.font.bold, italic=cell.font.italic,
                    color=cell.font.color,
                )

        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            max_len = max((len(str(c.value)) for c in col if c.value is not None), default=0)
            ws.column_dimensions[letter].width = min(max_len + 2, 50)

        ws[f'A{max_row + 2}'] = 'End of worksheet'
        ws[f'A{max_row + 2}'].font = Font(size=12)

    wb.save(filename)


def create_about_sheet_eeo4(filename, file_type='adjusted'):
    """Prepend an About sheet with a table of contents to an existing EEO-4 workbook."""
    sheet_descriptions = {
        'Category-Race-Sex':             'Job Category × Race/Ethnicity × Gender',
        'Type-Function-Sex':             'Work Type × Government Function × Gender',
        'Type-Function-Race':            'Work Type × Government Function × Race/Ethnicity',
        'Work-Salary-Function-Race-Sex': 'Work Type × Salary Range × Government Function × Race/Ethnicity × Gender',
    }

    wb = load_workbook(filename)
    if 'About' in wb.sheetnames:
        del wb['About']
    ws = wb.create_sheet('About', 0)

    ws['A1'] = 'About'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')

    ws['A2'] = (
        'This data has been collected from EEO-4 reports submitted by Massachusetts employers '
        ' (state and local government entities with 100+ employees) under the Frances Perkins '
        'Workplace Equity Act of 2024. Several steps were taken to protect data and identity of '
        'respondents. These include physical and logical access controls, cryptographically secure '
        'multi-party computation to provide end-to-end data confidentiality, and differential privacy '
        'techniques. In particular, all counts shown in these tables are approximate. '
        ' More information about data security can be found in the methodology section of the '
        'full 2026 Massachusetts Workforce Data Report. A table of contents is below.'
    )
    ws['A2'].font = Font(size=12)
    ws['A2'].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 60

    ws['A3'] = 'For more information, visit mass.gov/Workforce2026'
    ws['A3'].font = Font(size=12)

    ws['A4'] = 'Table of Contents'
    ws['A4'].font = Font(bold=True, size=12)

    row = 5
    for sheet_name, description in sheet_descriptions.items():
        ws[f'A{row}'] = sheet_name
        ws[f'A{row}'].font = Font(color='0563C1', underline='single', size=12)
        ws[f'A{row}'].hyperlink = Hyperlink(ref=f'A{row}', location=f"'{sheet_name}'!A1")
        ws[f'B{row}'] = description
        ws[f'B{row}'].font = Font(size=11)
        row += 1

    ws[f'A{ws.max_row + 2}'] = 'End of worksheet'
    ws[f'A{ws.max_row + 2}'].font = Font(size=12)

    ws.column_dimensions['A'].width = 80
    ws.column_dimensions['B'].width = 50

    wb.save(filename)


# =============================================================================
# Step 1: Load raw tables, write raw and zeroized workbooks
# =============================================================================

wfrg_all = pd.read_csv(os.path.join(input_dir, "WFRG_all_hires.csv"))
wfrg_new = pd.read_csv(os.path.join(input_dir, "WFRG_new_hires.csv"))
wfrg_combo_raw = pd.concat([wfrg_all, wfrg_new], ignore_index=True)

raw_tables = [pd.read_csv(os.path.join(input_dir, f)) for f in ["JRG.csv", "TFG.csv", "TFR.csv"]]
raw_tables.append(wfrg_combo_raw)
for i in range(len(raw_tables)):
    raw_tables[i].rename(columns=RENAME_DICT, inplace=True)

dataframes_raw = {
    'Category-Race-Sex':             raw_tables[0],
    'Type-Function-Sex':             raw_tables[1],
    'Type-Function-Race':            raw_tables[2],
    'Work-Salary-Function-Race-Sex': raw_tables[3],
}

out_raw = os.path.join(output_dir, "MA_Workforce_2026_EEO4_raw.xlsx")
write_tables_with_titles(out_raw, dataframes_raw)
create_about_sheet_eeo4(out_raw, file_type='raw')
print(f"→ saved {out_raw}")

dataframes_zeroized = {
    name: df.assign(Count=df['Count'].clip(lower=0))
    for name, df in dataframes_raw.items()
}

out_zeroized = os.path.join(output_dir, "MA_Workforce_2026_EEO4_remove_all_zeroes.xlsx")
write_tables_with_titles(out_zeroized, dataframes_zeroized)
create_about_sheet_eeo4(out_zeroized, file_type='zeroized')
print(f"→ saved {out_zeroized}")


# =============================================================================
# Step 2: Adjust 3-way side tables
# =============================================================================

# Reads raw CSVs from disk (original column names from eeo4_dp.py).
# Output *_adj.csv files also keep original column names; RENAME_DICT is
# applied in Step 4 when loading for the final workbook.
THREE_WAY_CONFIGS = [
    {"file": "JRG.csv", "max_cell": 2, "max_1way": 0,  "max_2way": 10},
    {"file": "TFG.csv", "max_cell": 1, "max_1way": 11, "max_2way": 13},
    {"file": "TFR.csv", "max_cell": 2, "max_1way": 33, "max_2way": 33},
]

for config in THREE_WAY_CONFIGS:
    csv_path = os.path.join(input_dir, config["file"])
    adjust_3way_table(
        csv_path=csv_path,
        count_col='Count',
        use_l1=True,
        max_cell_change=config["max_cell"],
        max_1way_marginal_change=config["max_1way"],
        max_2way_marginal_change=config["max_2way"],
        output_csv=csv_path.replace(".csv", "_adj.csv"),
    )


# =============================================================================
# Step 3: Build, adjust, and expand the 4-way WFRG table
# =============================================================================

# Re-load and concat; apply RENAME_DICT so Work_Type/Salary_Range exist for combining
wfrg_all = pd.read_csv(os.path.join(input_dir, "WFRG_all_hires.csv"))
wfrg_new = pd.read_csv(os.path.join(input_dir, "WFRG_new_hires.csv"))
wfrg_combo = pd.concat([wfrg_all, wfrg_new], ignore_index=True)
wfrg_combo.rename(columns=RENAME_DICT, inplace=True)

# Combine Work_Type and Salary_Range into a single column to reduce to 4 dimensions.
# Prefix 'Z' on new-hire rows so sorted order keeps them at the bottom of the table.
wfrg_combo['Work_Salary'] = wfrg_combo['Work_Type'] + '|' + wfrg_combo['Salary_Range']
wfrg_4way = wfrg_combo[['Work_Salary', 'Government_Function', 'Race_Ethnicity', 'Sex', 'Count']].copy()
wfrg_4way['Work_Salary'] = wfrg_4way['Work_Salary'].replace({'NEW HIRES|-': 'ZNEW HIRES|-'})
wfrg_4way.to_csv(os.path.join(input_dir, "WFRG_combo_4way.csv"), index=False)

adjust_4way_table(
    csv_path=os.path.join(input_dir, "WFRG_combo_4way.csv"),
    count_col='Count',
    use_l1=True,
    max_cell_change=3,
    max_1way_marginal_change=3,
    max_2way_marginal_change=11,
    max_3way_marginal_change=27,
    max_0way_marginal_change=0,
    output_csv=os.path.join(input_dir, "WFRG_combo_4way_adj.csv"),
)

# Expand the adjusted table back to 5 columns
wfrg_adj = pd.read_csv(os.path.join(input_dir, "WFRG_combo_4way_adj.csv"))
wfrg_adj['Work_Salary'] = wfrg_adj['Work_Salary'].replace({'ZNEW HIRES|-': 'NEW HIRES|-'})
wfrg_adj[['Work_Type', 'Salary_Range']] = wfrg_adj['Work_Salary'].str.split('|', expand=True)
wfrg_adj = wfrg_adj[['Work_Type', 'Salary_Range', 'Government_Function', 'Race_Ethnicity', 'Sex', 'Count']]
wfrg_adj.to_csv(os.path.join(input_dir, "WFRG_combo_adj.csv"), index=False)
print("→ saved WFRG_combo_adj.csv")


# =============================================================================
# Step 4: Write adjusted workbook
# =============================================================================

# JRG_adj/TFG_adj/TFR_adj have original column names → apply RENAME_DICT.
# WFRG_combo_adj already uses renamed columns → RENAME_DICT is a no-op for it.
dfs_adj = []
for filename in ["JRG_adj.csv", "TFG_adj.csv", "TFR_adj.csv", "WFRG_combo_adj.csv"]:
    df = pd.read_csv(os.path.join(input_dir, filename))
    df.rename(columns=RENAME_DICT, inplace=True)
    dfs_adj.append(df)

dataframes_adj = {
    'Category-Race-Sex':             dfs_adj[0],
    'Type-Function-Sex':             dfs_adj[1],
    'Type-Function-Race':            dfs_adj[2],
    'Work-Salary-Function-Race-Sex': dfs_adj[3],
}

out_adj = os.path.join(output_dir, "MA_Workforce_2026_EEO4_remove_structured_zeroes.xlsx")
write_tables_with_titles(out_adj, dataframes_adj)
create_about_sheet_eeo4(out_adj, file_type='adjusted')
print(f"→ saved {out_adj}")
